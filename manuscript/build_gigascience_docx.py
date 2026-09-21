#!/usr/bin/env python3
"""Build the review-ready GigaScience manuscript from the versioned Markdown.

The generated DOCX is a submission draft. Scientific text remains sourced from
GIGASCIENCE_RESEARCH_ARTICLE.md; the script adds journal-facing front matter,
figure legends, embedded deterministic figures, line numbers and page numbers.
"""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript" / "GIGASCIENCE_RESEARCH_ARTICLE.md"
OUTPUT = ROOT / "manuscript" / "RetinaCellTwin_GigaScience_Submission_Draft.docx"
FIGURES = ROOT / "figures"

TITLE = (
    "RetinaCellTwin Reproducible Multimodal and Cross Mission Audit of Retinal "
    "Responses to Spaceflight Artificial Gravity and Ground Analogues"
)

FIGURE_CAPTIONS = [
    (
        "Figure 1",
        "Evidence architecture. OSD-255 defines the frozen discovery panel; OSD-758 "
        "provides the independent artificial-gravity audit; OSD-203 provides time-resolved "
        "ground-analogue comparisons; OSD-557 and OSD-568 provide unpaired imaging context. "
        "Observed, derived and hypothetical layers remain separated.",
        "figure1_evidence_architecture.svg",
    ),
    (
        "Figure 2",
        "Cross-mission directional concordance. The OSD-255 signature does not reproduce "
        "as a universal gravity signature in OSD-758. The seven-day radiation comparison "
        "shows the only panel-enriched positive concordance and does not persist at later times.",
        "figure2_cross_mission_direction.svg",
    ),
    (
        "Figure 3",
        "Effect-level technical audits. Spearman correlations compare NASA differential-expression "
        "effects with contrasts recomputed from public variance-stabilized matrices and with effects "
        "recomputed after independent median-of-ratios normalization of public STAR counts.",
        "figure3_audit_concordance.svg",
    ),
    (
        "Figure 4",
        "Gene Ontology Biological Process context. The OSD-255 discovery panel reproduces known "
        "RR-9 biological context, whereas neither the seven-day radiation-concordant subset nor the "
        "nine-candidate subset contains a term passing the prespecified FDR boundary against the panel.",
        "figure4_go_context.svg",
    ),
    (
        "Figure 5",
        "Canonical retinal cell-class anchors across four bulk RNA-seq contrasts. Only the cone "
        "anchor Arr3 passes FDR in OSD-255; no anchor passes FDR in the independent or analogue "
        "contrasts. The audit provides context and is not a deconvolution analysis.",
        "figure5_cell_class_context.svg",
    ),
    (
        "Figure 6",
        "Unpaired cross-modal evidence matrix. Study-level contextual convergence is limited to "
        "photoreceptor integrity. Apoptosis has imaging-only FDR-supported evidence because no direct "
        "transcriptomic apoptosis test was prespecified. Oxidative-stress and vascular or barrier "
        "imaging domains do not pass the seven-endpoint Welch FDR boundary.",
        "figure6_cross_modal_context.svg",
    ),
]


def raw_page_break() -> str:
    return "```{=openxml}\n<w:p><w:r><w:br w:type=\"page\"/></w:r></w:p>\n```"


def clean_body(source: str) -> str:
    lines = source.splitlines()
    abstract_index = lines.index("## Abstract")
    body = "\n".join(lines[abstract_index:])
    heading_replacements = {
        "### RR-9 does not define a stable universal gravity signature":
            "### The RR 9 Signature Does Not Define a Stable Universal Gravity Signature",
        "### Sample-level sensitivity audit": "### Sample Level Sensitivity Audit",
        "### Count-level effect audit": "### Count Level Effect Audit",
        "### Retinal cell-class anchor context": "### Retinal Cell Class Anchor Context",
        "### Cross-modal convergence is confined to study-level photoreceptor context":
            "### Cross Modal Convergence Is Confined to Study Level Photoreceptor Context",
        "### Unpaired cross-modal synthesis": "### Unpaired Cross Modal Synthesis",
    }
    for old, new in heading_replacements.items():
        body = body.replace(old, new)
    return body


def build_markdown(temp_dir: Path) -> Path:
    body = clean_body(SOURCE.read_text(encoding="utf-8"))
    word_count = len(re.findall(r"\b[\w'-]+\b", body))
    front = f"""---
title: '{TITLE}'
---

**Research Article**

**Data Driven Multicellular Systems Biology**

Lissandra Kruse Fuganti^1^

^1^ Universidade Estadual de Ponta Grossa, Ponta Grossa, Parana, Brazil

**Corresponding author** Lissandra Kruse Fuganti, Universidade Estadual de Ponta Grossa, Ponta Grossa, Parana, Brazil. Email: t7426541@gmail.com. ORCID: 0009-0008-8189-112X.

**Manuscript word count** {word_count}

{raw_page_break()}
"""
    legends = ["\n# Figure Legends\n"]
    for label, caption, _ in FIGURE_CAPTIONS:
        legends.append(f"**{label}** {caption}\n")

    figure_pages: list[str] = []
    for label, caption, svg_name in FIGURE_CAPTIONS:
        png_name = svg_name.replace(".svg", ".png")
        png_path = temp_dir / png_name
        subprocess.run(
            [
                "inkscape",
                str(FIGURES / svg_name),
                "--export-type=png",
                f"--export-filename={png_path}",
                "--export-width=2400",
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        figure_pages.extend(
            [
                raw_page_break(),
                f"## {label}",
                "",
                f"![]({png_path.as_posix()}){{width=6.35in}}",
                "",
                f"*{caption}*",
                "",
            ]
        )

    combined = front + "\n" + body + "\n" + "\n".join(legends + figure_pages)
    markdown_path = temp_dir / "submission.md"
    markdown_path.write_text(combined, encoding="utf-8")
    return markdown_path


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(24)
    title.paragraph_format.keep_with_next = True

    for body_name in ("Body Text", "First Paragraph", "Compact", "Abstract"):
        if body_name not in styles:
            continue
        body_style = styles[body_name]
        body_style.font.name = "Times New Roman"
        body_style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        body_style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
        body_style.font.size = Pt(11.5)
        body_style.font.color.rgb = RGBColor(0, 0, 0)
        body_style.paragraph_format.line_spacing = 1.45
        body_style.paragraph_format.space_after = Pt(6)
        body_style.paragraph_format.widow_control = True

    heading_sizes = {"Heading 1": 15, "Heading 2": 12.5, "Heading 3": 11.5}
    for name, size in heading_sizes.items():
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(5)

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.line_spacing = 1.15
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_together = True


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, "Arial", 9)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)

        sect_pr = section._sectPr
        existing = sect_pr.find(qn("w:lnNumType"))
        if existing is None:
            existing = OxmlElement("w:lnNumType")
            sect_pr.append(existing)
        existing.set(qn("w:countBy"), "1")
        existing.set(qn("w:restart"), "continuous")
        existing.set(qn("w:distance"), "360")

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("RetinaCellTwin cross mission retinal audit")
        set_run_font(hr, "Arial", 8.5)
        hr.font.color.rgb = RGBColor(89, 89, 89)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, "PAGE")


def style_content(doc: Document) -> None:
    in_references = False
    after_title = False
    figure_heading_seen = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name

        if style_name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            after_title = True
        elif after_title and text and style_name == "Normal":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.15
            if text.startswith("Research Article") or text.startswith("Data Driven"):
                for run in paragraph.runs:
                    set_run_font(run, "Arial", 10.5, True)
            elif text.startswith("Corresponding author") or text.startswith("Manuscript word count"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                after_title = False

        if style_name == "Heading 1":
            in_references = text == "References"
            if text == "Figure Legends":
                in_references = False
        elif style_name == "Heading 2" and text.startswith("Figure "):
            figure_heading_seen = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif style_name == "Normal" and in_references and re.match(r"^\d+\.\s", text):
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.line_spacing = 1.15
            for run in paragraph.runs:
                set_run_font(run, "Times New Roman", 10.5)
        elif figure_heading_seen and style_name == "Normal" and text:
            if paragraph._p.xpath(".//a:blip"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
            elif text.startswith("Evidence architecture") or text.startswith("Cross-mission") or text.startswith("Effect-level") or text.startswith("Gene Ontology") or text.startswith("Canonical retinal") or text.startswith("Unpaired cross-modal"):
                paragraph.style = doc.styles["Figure Caption"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                figure_heading_seen = False

        if text.startswith("Figure ") and style_name == "Normal" and paragraph.runs:
            paragraph.paragraph_format.keep_together = True

    for table in doc.tables:
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.1
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        set_run_font(run, "Times New Roman", 9.5, row_index == 0)


def sanitize_image_metadata(doc: Document) -> None:
    """Replace Pandoc's temporary image paths with stable accessible labels."""
    if len(doc.inline_shapes) != len(FIGURE_CAPTIONS):
        raise RuntimeError(
            f"Expected {len(FIGURE_CAPTIONS)} embedded figures, "
            f"found {len(doc.inline_shapes)}"
        )
    for shape, (label, caption, _) in zip(doc.inline_shapes, FIGURE_CAPTIONS):
        inline = shape._inline
        inline.docPr.set("title", label)
        inline.docPr.set("descr", caption)
        picture_properties = inline.graphic.graphicData.pic.nvPicPr.cNvPr
        picture_properties.set("name", label)
        picture_properties.set("descr", caption)


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="retinacelltwin_docx_") as temp_name:
        temp_dir = Path(temp_name)
        markdown_path = build_markdown(temp_dir)
        raw_docx = temp_dir / "raw.docx"
        subprocess.run(
            [
                "pandoc",
                str(markdown_path),
                "--from=markdown+raw_attribute",
                "--to=docx",
                f"--resource-path={temp_dir}:{ROOT}",
                "--output",
                str(raw_docx),
            ],
            check=True,
        )

        doc = Document(raw_docx)
        configure_styles(doc)
        configure_sections(doc)
        style_content(doc)
        sanitize_image_metadata(doc)

        core = doc.core_properties
        core.title = TITLE
        core.subject = "GigaScience Research Article submission draft"
        core.author = "Lissandra Kruse Fuganti"
        core.keywords = "retina, spaceflight, NASA OSDR, transcriptomics, microscopy, reproducibility"
        core.comments = "Generated reproducibly from the versioned RetinaCellTwin manuscript and figures."

        doc.save(OUTPUT)
        print(OUTPUT)


if __name__ == "__main__":
    main()
